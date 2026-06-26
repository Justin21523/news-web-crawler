from __future__ import annotations

import html
import json
import uuid
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from app.core.config import Settings, get_settings
from app.schemas.ml import (
    MLDiagnosticsReportBulkResponse,
    MLDiagnosticsReportCompareResponse,
    MLDiagnosticsReportDetailResponse,
    MLDiagnosticsReportListResponse,
    MLDiagnosticsReportRecord,
    MLDiagnosticsReportMetadataUpdate,
    MLModelName,
    MLTarget,
)
from app.services.database import Database
from app.services.ml import MLService


REPORT_SECTIONS = [
    "cover",
    "executive_summary",
    "metrics",
    "confusion_matrix",
    "class_quality",
    "failure_narrative",
    "recommendations",
    "error_samples",
    "artifact_comparison",
    "article_explanations",
    "appendix_filters",
    "appendix_predictions",
]

DEFAULT_REPORT_SECTIONS = [
    "cover",
    "executive_summary",
    "metrics",
    "confusion_matrix",
    "class_quality",
    "failure_narrative",
    "recommendations",
    "error_samples",
    "artifact_comparison",
    "article_explanations",
    "appendix_filters",
    "appendix_predictions",
]


class MLDiagnosticsReportService:
    def __init__(self, db: Database | None = None, settings: Settings | None = None):
        self.db = db or Database()
        self.settings = settings or get_settings()
        self.ml = MLService(self.db, self.settings)
        self.reports_dir = self.settings.resolved_reports_dir / "ml_diagnostics"
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        self.bulk_exports_dir = self.reports_dir / "_bulk_exports"
        self.bulk_exports_dir.mkdir(parents=True, exist_ok=True)
        self._ensure_metadata_table()

    def generate(
        self,
        target: MLTarget = "source",
        model: MLModelName = "logistic_regression",
        job_id: int | None = None,
        artifact_ids: list[str] | None = None,
        article_ids: list[str] | None = None,
        explanation_modes: list[str] | None = None,
        template: str = "portfolio",
        sections: list[str] | None = None,
        report_title: str = "ML Diagnostics Report",
        prepared_for: str | None = None,
        **params: Any,
    ) -> MLDiagnosticsReportRecord:
        report_id = f"ml-diagnostics-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}"
        report_dir = self.reports_dir / report_id
        report_dir.mkdir(parents=True, exist_ok=True)
        artifact_ids = [item for item in (artifact_ids or []) if item]
        article_ids = [item for item in (article_ids or []) if item]
        explanation_modes = [item for item in (explanation_modes or ["linear_coefficients", "tree_path"]) if item]
        selected_sections, section_notes = self._sections(sections)
        template = template or "portfolio"
        report_title = report_title or "ML Diagnostics Report"
        base_params = self._clean_params(params)
        base_params.update({"target": target, "model": model})

        diagnostics = self.ml.diagnostics(target=target, model=model, **self._fit_params(base_params)).model_dump()
        errors = self.ml.error_samples(
            target=target,
            model=model,
            actual=base_params.get("actual"),
            predicted=base_params.get("predicted"),
            class_label=base_params.get("class_label"),
            page_size=10000,
            **self._fit_params(base_params),
        ).model_dump()
        comparison = self.ml.compare_artifact_diagnostics(target, artifact_ids).model_dump() if artifact_ids else {}
        explanations = self._explanations(article_ids, target, explanation_modes, base_params)
        narrative = self._narrative(diagnostics, comparison, errors.get("items", []))
        section_status = self._section_status(selected_sections, diagnostics, comparison, errors.get("items", []), explanations)
        payload = {
            "report_id": report_id,
            "job_id": job_id,
            "created_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "template": template,
            "sections": selected_sections,
            "section_status": section_status,
            "report_title": report_title,
            "prepared_for": prepared_for,
            "params": {**base_params, "artifact_ids": artifact_ids, "article_ids": article_ids, "explanation_modes": explanation_modes},
            "diagnostics": diagnostics,
            "error_samples": errors.get("items", []),
            "all_predictions": diagnostics.get("sample_predictions", []),
            "artifact_comparison": comparison,
            "explanations": explanations,
            "narrative": narrative,
            "notes": [*section_notes, *diagnostics.get("notes", []), *comparison.get("notes", [])] if comparison else [*section_notes, *diagnostics.get("notes", [])],
        }
        row_counts = {
            "error_samples": len(payload["error_samples"]),
            "all_predictions": len(payload["all_predictions"]),
            "artifact_comparison": len(comparison.get("items", [])) if comparison else 0,
            "explanations": len(explanations),
        }
        summary_metrics = diagnostics.get("metrics", {})

        (report_dir / "report.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        (report_dir / "report.md").write_text(self._markdown(payload), encoding="utf-8")
        (report_dir / "report.html").write_text(self._html(payload), encoding="utf-8")
        self._excel(report_dir / "report.xlsx", payload)
        manifest = {
            "report_id": report_id,
            "job_id": job_id,
            "created_at": payload["created_at"],
            "template": template,
            "sections": selected_sections,
            "section_status": section_status,
            "report_title": report_title,
            "prepared_for": prepared_for,
            "preview_url": f"/ml/diagnostics/reports/{report_id}/download?format=html",
            "print_url": f"/ml/diagnostics/reports/{report_id}/download?format=html",
            "pdf_url": f"/ml/diagnostics/reports/{report_id}/download?format=pdf",
            "pdf_status": "not_generated",
            "pdf_fallback_reason": None,
            "pdf_generated_at": None,
            "params": payload["params"],
            "files": ["manifest.json", "report.html", "report.md", "report.xlsx", "report.json"],
            "summary_metrics": summary_metrics,
            "row_counts": row_counts,
            "notes": payload["notes"],
        }
        (report_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return MLDiagnosticsReportRecord(**manifest)

    def list_reports(
        self,
        limit: int = 50,
        q: str | None = None,
        target: str | None = None,
        model: str | None = None,
        template: str | None = None,
        pdf_status: str | None = None,
        status: str | None = "active",
        tag: str | None = None,
        include_trashed: bool = False,
        date_from: str | None = None,
        date_to: str | None = None,
        sort_by: str = "created_at",
        sort_order: str = "desc",
        page: int = 1,
        page_size: int | None = None,
    ) -> MLDiagnosticsReportListResponse:
        records = [MLDiagnosticsReportRecord(**self._normalize_manifest(self._read_json(path, {}))) for path in self.reports_dir.glob("*/manifest.json")]
        filtered = [
            record
            for record in records
            if self._matches_report(record, q, target, model, template, pdf_status, status, tag, include_trashed, date_from, date_to)
        ]
        filtered.sort(key=lambda record: self._sort_value(record, sort_by), reverse=sort_order != "asc")
        effective_page_size = page_size or limit
        start = (max(page, 1) - 1) * max(effective_page_size, 1)
        end = start + min(max(effective_page_size, 1), 100)
        return MLDiagnosticsReportListResponse(items=filtered[start:end], total=len(filtered))

    def detail(self, report_id: str) -> MLDiagnosticsReportDetailResponse | None:
        report_dir = self._report_dir(report_id)
        manifest = self._normalize_manifest(self._read_json(report_dir / "manifest.json", {}))
        if not manifest:
            return None
        payload = self._read_json(report_dir / "report.json", {})
        return MLDiagnosticsReportDetailResponse(**manifest, artifact_comparison=payload.get("artifact_comparison"), diagnostics=payload.get("diagnostics"))

    def update_metadata(self, report_id: str, update: MLDiagnosticsReportMetadataUpdate) -> MLDiagnosticsReportRecord | None:
        if not (self._report_dir(report_id) / "manifest.json").exists():
            return None
        current = self._metadata_map([report_id]).get(report_id, self._default_metadata())
        status = update.status or current["status"]
        tags = self._clean_tags(update.tags if update.tags is not None else current["tags"])
        self._write_metadata(report_id, status, tags)
        detail = self.detail(report_id)
        return MLDiagnosticsReportRecord(**detail.model_dump()) if detail else None

    def bulk_update(self, action: str, report_ids: list[str], tags: list[str] | None = None) -> MLDiagnosticsReportBulkResponse:
        clean_ids = [item for item in dict.fromkeys(report_ids) if item][:100]
        clean_tags = self._clean_tags(tags or [])
        missing = []
        items = []
        for report_id in clean_ids:
            if not (self._report_dir(report_id) / "manifest.json").exists():
                missing.append(report_id)
                continue
            current = self._metadata_map([report_id]).get(report_id, self._default_metadata())
            status = current["status"]
            next_tags = list(current["tags"])
            if action == "archive":
                status = "archived"
            elif action == "trash":
                status = "trashed"
            elif action == "restore":
                status = "active"
            elif action == "add_tags":
                next_tags = self._clean_tags([*next_tags, *clean_tags])
            elif action == "remove_tags":
                remove = set(clean_tags)
                next_tags = [tag for tag in next_tags if tag not in remove]
            elif action == "set_tags":
                next_tags = clean_tags
            else:
                raise ValueError("Unsupported report bulk action")
            self._write_metadata(report_id, status, next_tags)
            record = self.detail(report_id)
            if record:
                items.append(MLDiagnosticsReportRecord(**record.model_dump()))
        return MLDiagnosticsReportBulkResponse(action=action, updated=len(items), missing=missing, items=items)

    def bulk_export(self, report_ids: list[str], formats: list[str]) -> tuple[Path, str, str, dict[str, str]]:
        clean_ids = [item for item in dict.fromkeys(report_ids) if item][:100]
        clean_formats = [item for item in (formats or ["html", "pdf", "xlsx", "json", "manifest"]) if item in {"html", "pdf", "xlsx", "json", "markdown", "manifest"}]
        if not clean_ids:
            raise ValueError("No report IDs supplied for bulk export")
        export_id = f"report-bulk-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}"
        zip_path = self.bulk_exports_dir / f"{export_id}.zip"
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for report_id in clean_ids:
                report_dir = self._report_dir(report_id)
                if not (report_dir / "manifest.json").exists():
                    continue
                record = self.detail(report_id)
                if record:
                    archive.writestr(f"{report_id}/metadata.json", json.dumps(record.model_dump(), ensure_ascii=False, indent=2))
                for fmt in clean_formats:
                    try:
                        if fmt == "manifest":
                            path = report_dir / "manifest.json"
                            filename = "manifest.json"
                        else:
                            path, media_type, filename, headers = self.download(report_id, fmt)
                            if fmt == "pdf" and headers.get("X-Report-Fallback") == "html":
                                filename = "PDF_FALLBACK.html"
                        if path.exists():
                            archive.write(path, f"{report_id}/{filename}")
                    except (FileNotFoundError, ValueError):
                        continue
        return zip_path, "application/zip", f"{export_id}.zip", {}

    def compare_reports(self, report_ids: list[str]) -> MLDiagnosticsReportCompareResponse:
        selected_ids = [item for item in report_ids[:5] if item]
        records = []
        payloads = {}
        notes = []
        for report_id in selected_ids:
            report_dir = self._report_dir(report_id)
            manifest = self._normalize_manifest(self._read_json(report_dir / "manifest.json", {}))
            if not manifest:
                notes.append(f"Report not found: {report_id}")
                continue
            records.append(MLDiagnosticsReportRecord(**manifest))
            payloads[report_id] = self._read_json(report_dir / "report.json", {})
        baseline = records[0] if records else None
        metrics = [{"report_id": record.report_id, **record.summary_metrics} for record in records]
        metric_deltas = []
        row_count_deltas = []
        if baseline:
            baseline_metrics = baseline.summary_metrics
            baseline_rows = baseline.row_counts
            metric_keys = sorted({key for record in records for key in record.summary_metrics})
            row_keys = sorted({key for record in records for key in record.row_counts})
            for record in records:
                for key in metric_keys:
                    value = self._numeric(record.summary_metrics.get(key))
                    base = self._numeric(baseline_metrics.get(key))
                    metric_deltas.append({"report_id": record.report_id, "metric": key, "value": value, "baseline": base, "delta": None if value is None or base is None else round(value - base, 6)})
                for key in row_keys:
                    value = self._numeric(record.row_counts.get(key))
                    base = self._numeric(baseline_rows.get(key))
                    row_count_deltas.append({"report_id": record.report_id, "name": key, "value": value, "baseline": base, "delta": None if value is None or base is None else value - base})
        section_names = sorted({section for record in records for section in record.sections})
        section_matrix = [
            {"section": section, **{record.report_id: record.section_status.get(section, "missing") if section in record.sections else "missing" for record in records}}
            for section in section_names
        ]
        failure_summary = []
        for record in records:
            diagnostics = payloads.get(record.report_id, {}).get("diagnostics", {})
            narrative = payloads.get(record.report_id, {}).get("narrative", {})
            failure_summary.append({
                "report_id": record.report_id,
                "failure_factors": len(diagnostics.get("failure_factors", [])),
                "recommendations": len(diagnostics.get("recommendations", [])),
                "summary": narrative.get("executive_summary", ""),
            })
        return MLDiagnosticsReportCompareResponse(
            baseline_report_id=baseline.report_id if baseline else None,
            items=records,
            metrics=metrics,
            metric_deltas=metric_deltas,
            row_count_deltas=row_count_deltas,
            section_matrix=section_matrix,
            failure_summary=failure_summary,
            notes=notes,
        )

    def compare_download(self, report_ids: list[str], fmt: str) -> tuple[Path, str, str, dict[str, str]]:
        comparison = self.compare_reports(report_ids)
        if not comparison.items:
            raise FileNotFoundError("No diagnostics reports found for comparison")
        export_id = f"report-compare-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}-{uuid.uuid4().hex[:8]}"
        html_path = self.bulk_exports_dir / f"{export_id}.html"
        pdf_path = self.bulk_exports_dir / f"{export_id}.pdf"
        html_path.write_text(self._compare_html(comparison), encoding="utf-8")
        if fmt == "html":
            return html_path, "text/html; charset=utf-8", html_path.name, {}
        if fmt != "pdf":
            raise ValueError("Unsupported comparison export format")
        try:
            self._render_pdf(html_path, pdf_path)
        except Exception as exc:
            reason = self._short_error(exc)
            return html_path, "text/html; charset=utf-8", html_path.name, {"X-Report-Fallback": "html", "X-Report-Fallback-Reason": reason[:180]}
        return pdf_path, "application/pdf", pdf_path.name, {}

    def download(self, report_id: str, fmt: str) -> tuple[Path, str, str, dict[str, str]]:
        mapping = {
            "html": ("report.html", "text/html; charset=utf-8"),
            "markdown": ("report.md", "text/markdown; charset=utf-8"),
            "xlsx": ("report.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "json": ("report.json", "application/json; charset=utf-8"),
        }
        if fmt == "pdf":
            return self._pdf_or_html(report_id)
        if fmt not in mapping:
            raise ValueError("Unsupported diagnostics report format")
        filename, media_type = mapping[fmt]
        path = self._report_dir(report_id) / filename
        if not path.exists():
            raise FileNotFoundError("Diagnostics report file not found")
        return path, media_type, filename, {}

    def _pdf_or_html(self, report_id: str) -> tuple[Path, str, str, dict[str, str]]:
        report_dir = self._report_dir(report_id)
        html_path = report_dir / "report.html"
        pdf_path = report_dir / "report.pdf"
        manifest_path = report_dir / "manifest.json"
        if not html_path.exists():
            raise FileNotFoundError("Diagnostics report HTML not found")
        manifest = self._normalize_manifest(self._read_json(manifest_path, {}))
        if pdf_path.exists():
            return pdf_path, "application/pdf", "report.pdf", {}
        try:
            self._render_pdf(html_path, pdf_path)
        except Exception as exc:  # optional dependency and browser installation can fail in several ways
            reason = self._short_error(exc)
            manifest.update({
                "pdf_status": "fallback_html",
                "pdf_fallback_reason": reason,
                "pdf_generated_at": None,
            })
            manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            return html_path, "text/html; charset=utf-8", "report.html", {"X-Report-Fallback": "html", "X-Report-Fallback-Reason": reason[:180]}
        manifest.update({
            "pdf_status": "ready",
            "pdf_fallback_reason": None,
            "pdf_generated_at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "files": sorted(set([*manifest.get("files", []), "report.pdf"])),
        })
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        return pdf_path, "application/pdf", "report.pdf", {}

    def _render_pdf(self, html_path: Path, pdf_path: Path) -> None:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page()
            page.goto(html_path.resolve().as_uri(), wait_until="networkidle")
            page.pdf(path=str(pdf_path), format="A4", print_background=True, prefer_css_page_size=True)
            browser.close()

    def _normalize_manifest(self, manifest: dict[str, Any]) -> dict[str, Any]:
        if not manifest:
            return {}
        report_id = manifest.get("report_id", "")
        manifest.setdefault("preview_url", f"/ml/diagnostics/reports/{report_id}/download?format=html")
        manifest.setdefault("print_url", f"/ml/diagnostics/reports/{report_id}/download?format=html")
        manifest.setdefault("pdf_url", f"/ml/diagnostics/reports/{report_id}/download?format=pdf")
        manifest.setdefault("pdf_status", "ready" if (self._report_dir(report_id) / "report.pdf").exists() else "not_generated")
        manifest.setdefault("pdf_fallback_reason", None)
        manifest.setdefault("pdf_generated_at", None)
        manifest.setdefault("section_status", {})
        manifest.setdefault("summary_metrics", {})
        manifest.setdefault("row_counts", {})
        manifest.setdefault("notes", [])
        metadata = self._metadata_map([report_id]).get(report_id, self._default_metadata())
        manifest["status"] = metadata["status"]
        manifest["tags"] = metadata["tags"]
        manifest["archived_at"] = metadata["archived_at"]
        manifest["trashed_at"] = metadata["trashed_at"]
        manifest["metadata_updated_at"] = metadata["metadata_updated_at"]
        return manifest

    def _ensure_metadata_table(self) -> None:
        self.db.execute(
            """
            CREATE TABLE IF NOT EXISTS report_metadata (
                report_id TEXT PRIMARY KEY,
                status TEXT NOT NULL DEFAULT 'active',
                tags_json TEXT NOT NULL DEFAULT '[]',
                archived_at TEXT,
                trashed_at TEXT,
                updated_at TEXT DEFAULT (datetime('now'))
            )
            """
        )
        self.db.execute("CREATE INDEX IF NOT EXISTS idx_report_metadata_status ON report_metadata(status)")
        self.db.execute("CREATE INDEX IF NOT EXISTS idx_report_metadata_updated_at ON report_metadata(updated_at)")

    def _metadata_map(self, report_ids: list[str]) -> dict[str, dict[str, Any]]:
        if not report_ids:
            return {}
        placeholders = ",".join("?" for _ in report_ids)
        rows = self.db.fetch_all(f"SELECT * FROM report_metadata WHERE report_id IN ({placeholders})", tuple(report_ids))
        output = {}
        for row in rows:
            output[str(row["report_id"])] = {
                "status": row["status"] or "active",
                "tags": self._clean_tags(json.loads(row["tags_json"] or "[]")),
                "archived_at": row["archived_at"],
                "trashed_at": row["trashed_at"],
                "metadata_updated_at": row["updated_at"],
            }
        return output

    def _default_metadata(self) -> dict[str, Any]:
        return {"status": "active", "tags": [], "archived_at": None, "trashed_at": None, "metadata_updated_at": None}

    def _write_metadata(self, report_id: str, status: str, tags: list[str]) -> None:
        archived_at = datetime.utcnow().isoformat(timespec="seconds") + "Z" if status == "archived" else None
        trashed_at = datetime.utcnow().isoformat(timespec="seconds") + "Z" if status == "trashed" else None
        self.db.execute(
            """
            INSERT INTO report_metadata(report_id, status, tags_json, archived_at, trashed_at, updated_at)
            VALUES(?, ?, ?, ?, ?, datetime('now'))
            ON CONFLICT(report_id) DO UPDATE SET
                status=excluded.status,
                tags_json=excluded.tags_json,
                archived_at=excluded.archived_at,
                trashed_at=excluded.trashed_at,
                updated_at=datetime('now')
            """,
            (report_id, status, json.dumps(self._clean_tags(tags), ensure_ascii=False), archived_at, trashed_at),
        )

    def _clean_tags(self, tags: list[Any]) -> list[str]:
        output = []
        for tag in tags:
            value = str(tag).strip().lower()
            if value and value not in output:
                output.append(value[:40])
        return output[:20]

    def _matches_report(
        self,
        record: MLDiagnosticsReportRecord,
        q: str | None,
        target: str | None,
        model: str | None,
        template: str | None,
        pdf_status: str | None,
        status: str | None,
        tag: str | None,
        include_trashed: bool,
        date_from: str | None,
        date_to: str | None,
    ) -> bool:
        params = record.params or {}
        haystack = " ".join(str(value or "") for value in [record.report_id, record.report_title, record.prepared_for, record.template, params.get("target"), params.get("model")]).lower()
        if q and q.lower() not in haystack:
            return False
        if target and str(params.get("target")) != target:
            return False
        if model and str(params.get("model")) != model:
            return False
        if template and record.template != template:
            return False
        if pdf_status and record.pdf_status != pdf_status:
            return False
        if not include_trashed and record.status == "trashed":
            return False
        if status and record.status != status:
            return False
        if tag and tag.lower() not in record.tags:
            return False
        created = record.created_at or ""
        if date_from and created[:10] < date_from:
            return False
        if date_to and created[:10] > date_to:
            return False
        return True

    def _sort_value(self, record: MLDiagnosticsReportRecord, sort_by: str) -> Any:
        if sort_by == "report_title":
            return record.report_title.lower()
        if sort_by == "accuracy":
            return self._numeric(record.summary_metrics.get("accuracy")) or -1
        if sort_by == "f1_macro":
            return self._numeric(record.summary_metrics.get("f1_macro")) or -1
        if sort_by == "error_samples":
            return record.row_counts.get("error_samples", 0)
        return record.created_at or ""

    def _numeric(self, value: Any) -> float | None:
        try:
            if value is None or value == "":
                return None
            return float(value)
        except (TypeError, ValueError):
            return None

    def _short_error(self, exc: Exception) -> str:
        message = str(exc).replace("\n", " ").strip()
        if not message:
            message = exc.__class__.__name__
        if "playwright" in message.lower() or exc.__class__.__name__ == "ModuleNotFoundError":
            return "Optional Playwright dependency or Chromium browser is not installed."
        return message[:240]

    def _explanations(self, article_ids: list[str], target: MLTarget, modes: list[str], params: dict[str, Any]) -> list[dict[str, Any]]:
        if not article_ids:
            rows = self.db.fetch_all("SELECT article_id FROM articles ORDER BY COALESCE(publish_date, crawled_at) DESC LIMIT 3")
            article_ids = [str(row["article_id"]) for row in rows]
        output = []
        for article_id in article_ids[:20]:
            for mode in modes:
                result = self.ml.article_explanation(article_id=article_id, target=target, mode=mode, **self._fit_params(params))
                output.append(result.model_dump())
        return output

    def _fit_params(self, params: dict[str, Any]) -> dict[str, Any]:
        allowed = {"source", "category", "date_from", "date_to", "test_size", "feature_limit", "min_class_count", "collapse_rare", "max_depth", "min_samples_split", "min_samples_leaf", "criterion", "limit"}
        return {key: value for key, value in params.items() if key in allowed and value not in (None, "")}

    def _clean_params(self, params: dict[str, Any]) -> dict[str, Any]:
        output = {}
        for key, value in params.items():
            if value not in (None, ""):
                output[key] = value
        return output

    def _sections(self, sections: list[str] | None) -> tuple[list[str], list[str]]:
        if not sections:
            return list(DEFAULT_REPORT_SECTIONS), []
        selected = []
        ignored = []
        for section in sections:
            value = section.strip()
            if value in REPORT_SECTIONS and value not in selected:
                selected.append(value)
            elif value:
                ignored.append(value)
        notes = [f"Ignored unsupported report sections: {', '.join(ignored)}."] if ignored else []
        return selected or list(DEFAULT_REPORT_SECTIONS), notes

    def _section_status(
        self,
        sections: list[str],
        diagnostics: dict[str, Any],
        comparison: dict[str, Any],
        error_samples: list[dict[str, Any]],
        explanations: list[dict[str, Any]],
    ) -> dict[str, str]:
        status = {}
        for section in sections:
            ready = True
            if section == "artifact_comparison":
                ready = bool(comparison.get("items") or comparison.get("metric_deltas"))
            elif section == "error_samples":
                ready = bool(error_samples)
            elif section == "article_explanations":
                ready = bool(explanations)
            elif section == "confusion_matrix":
                ready = bool(diagnostics.get("labels") and diagnostics.get("confusion_matrix"))
            elif section == "class_quality":
                ready = bool(diagnostics.get("class_metrics"))
            status[section] = "ready" if ready else "empty"
        return status

    def _narrative(self, diagnostics: dict[str, Any], comparison: dict[str, Any], error_samples: list[dict[str, Any]]) -> dict[str, Any]:
        metrics = diagnostics.get("metrics", {})
        factors = diagnostics.get("failure_factors", [])
        recommendations = diagnostics.get("recommendations", [])
        accuracy = metrics.get("accuracy")
        macro_f1 = metrics.get("f1_macro")
        summary = (
            f"The selected model reached accuracy {accuracy} and macro F1 {macro_f1}. "
            "Macro F1 is the primary portfolio signal because it exposes class imbalance and minority-class weakness."
        )
        if accuracy is None and macro_f1 is None:
            summary = "The selected model does not have enough evaluation output to produce a reliable performance summary."
        factor_lines = [
            f"{item.get('title', item.get('factor', 'Model risk'))}: {item.get('detail', 'Review the affected class or data slice.')}"
            for item in factors[:5]
        ] or ["No explicit failure factors were generated. Treat this as a caveat, not proof that the model is production-ready."]
        rec_lines = [
            f"{item.get('title', 'Recommended action')}: {item.get('detail', 'Review training data and diagnostics before deployment.')}"
            for item in recommendations[:5]
        ] or ["Review training data balance, label quality, and error samples before treating this model as production-ready."]
        return {
            "executive_summary": summary,
            "failure_factor_narrative": factor_lines,
            "recommendation_narrative": rec_lines,
            "comparison_summary": self._comparison_summary(comparison),
            "error_sample_summary": f"{len(error_samples)} filtered misclassified samples are included for drill-down review.",
        }

    def _comparison_summary(self, comparison: dict[str, Any]) -> list[str]:
        items = comparison.get("items", [])
        best = comparison.get("best_artifact")
        if not items:
            return ["No persisted artifacts were selected for comparison in this report."]
        lines = [f"{len(items)} persisted artifacts were compared."]
        if best:
            score = best.get("metrics", {}).get("f1_macro") or best.get("metrics", {}).get("accuracy")
            lines.append(f"Best artifact: {best.get('artifact_id')} using {best.get('model')} with score {score}.")
        for delta in comparison.get("metric_deltas", [])[:5]:
            lines.append(f"{delta.get('metric')}: best {delta.get('best_value')} vs worst {delta.get('worst_value')} across compared artifacts.")
        return lines

    def _markdown(self, payload: dict[str, Any]) -> str:
        diagnostics = payload["diagnostics"]
        narrative = payload.get("narrative", {})
        sections = payload.get("sections", DEFAULT_REPORT_SECTIONS)
        lines = [
            f"# {payload.get('report_title') or 'ML Diagnostics Report'}",
            "",
            f"- Report ID: `{payload['report_id']}`",
            f"- Created: {payload['created_at']}",
            f"- Target: `{payload['params'].get('target')}`",
            f"- Model: `{payload['params'].get('model')}`",
            f"- Template: `{payload.get('template', 'portfolio')}`",
        ]
        if payload.get("prepared_for"):
            lines.append(f"- Prepared for: {payload['prepared_for']}")
        for section in sections:
            if section == "cover":
                continue
            if section == "executive_summary":
                lines.extend(["", "## Executive Summary", narrative.get("executive_summary", "_No summary available_")])
            elif section == "metrics":
                lines.extend(["## Metrics", self._markdown_table([diagnostics.get("metrics", {})])])
            elif section == "confusion_matrix":
                lines.extend(["## Confusion Matrix", self._markdown_table(self._matrix_rows(diagnostics.get("labels", []), diagnostics.get("confusion_matrix", [])))])
            elif section == "class_quality":
                lines.extend(["## Class Quality", self._markdown_table(diagnostics.get("class_metrics", []))])
            elif section == "failure_narrative":
                lines.extend(["## Failure Factor Narrative", "\n".join(f"- {line}" for line in narrative.get("failure_factor_narrative", [])), self._markdown_table(diagnostics.get("failure_factors", []))])
            elif section == "recommendations":
                lines.extend(["## Recommendations", "\n".join(f"- {line}" for line in narrative.get("recommendation_narrative", [])), self._markdown_table(diagnostics.get("recommendations", []))])
            elif section == "error_samples":
                lines.extend(["## Error Samples", self._markdown_table(payload.get("error_samples", []))])
            elif section == "artifact_comparison":
                lines.extend(["## Model Version Comparison", "\n".join(f"- {line}" for line in narrative.get("comparison_summary", [])), self._markdown_table(payload.get("artifact_comparison", {}).get("metric_deltas", []))])
            elif section == "article_explanations":
                lines.extend(["## Article Explanations", self._markdown_table([{k: v for k, v in item.items() if k in {"article_id", "target", "mode", "prediction", "actual_label", "model", "explanation"}} for item in payload.get("explanations", [])])])
            elif section == "appendix_filters":
                lines.extend(["## Appendix: Filters", self._markdown_table([payload.get("params", {})])])
            elif section == "appendix_predictions":
                lines.extend(["## Appendix: Predictions", self._markdown_table(payload.get("all_predictions", [])[:100])])
        return "\n\n".join(lines) + "\n"

    def _html(self, payload: dict[str, Any]) -> str:
        diagnostics = payload["diagnostics"]
        narrative = payload.get("narrative", {})
        selected = payload.get("sections", DEFAULT_REPORT_SECTIONS)
        blocks = []
        for section in selected:
            if section == "cover":
                prepared_for = f"<p>Prepared for {html.escape(str(payload.get('prepared_for')))}</p>" if payload.get("prepared_for") else ""
                blocks.append(f"""<header class="cover report-section"><div class="eyebrow">{html.escape(str(payload.get('template', 'portfolio'))).title()} Report</div><h1>{html.escape(payload.get('report_title') or 'ML Diagnostics Report')}</h1><p>Target <strong>{html.escape(str(payload['params'].get('target')))}</strong> · Model <strong>{html.escape(str(payload['params'].get('model')))}</strong></p><p>Report <code>{html.escape(payload['report_id'])}</code> · {html.escape(payload['created_at'])}</p>{prepared_for}</header>""")
            elif section == "executive_summary":
                cards = "".join(f"<div class=\"metric\"><span>{html.escape(str(key))}</span><strong>{html.escape(str(value))}</strong></div>" for key, value in diagnostics.get("metrics", {}).items())
                blocks.append(f"<section class=\"report-section\"><h2>Executive Summary</h2><p>{html.escape(narrative.get('executive_summary', 'No summary available.'))}</p><div class=\"metrics\">{cards}</div></section>")
            elif section == "metrics":
                blocks.append(f"<section class=\"report-section\"><h2>Metrics</h2>{self._html_table([diagnostics.get('metrics', {})])}</section>")
            elif section == "confusion_matrix":
                blocks.append(f"<section class=\"report-section\"><h2>Confusion Matrix</h2>{self._html_table(self._matrix_rows(diagnostics.get('labels', []), diagnostics.get('confusion_matrix', [])))}</section>")
            elif section == "class_quality":
                blocks.append(f"<section class=\"report-section\"><h2>Class Quality</h2>{self._html_table(diagnostics.get('class_metrics', []))}</section>")
            elif section == "failure_narrative":
                blocks.append(f"<section class=\"report-section\"><h2>Failure Factor Narrative</h2>{self._html_list(narrative.get('failure_factor_narrative', []))}<h3>Failure Factor Table</h3>{self._html_table(diagnostics.get('failure_factors', []))}</section>")
            elif section == "recommendations":
                blocks.append(f"<section class=\"report-section\"><h2>Recommendations</h2>{self._html_list(narrative.get('recommendation_narrative', []))}{self._html_table(diagnostics.get('recommendations', []))}</section>")
            elif section == "error_samples":
                blocks.append(f"<section class=\"report-section\"><h2>Error Samples</h2><p>{html.escape(narrative.get('error_sample_summary', ''))}</p>{self._html_table(payload.get('error_samples', []))}</section>")
            elif section == "artifact_comparison":
                blocks.append(f"<section class=\"report-section\"><h2>Model Version Comparison</h2>{self._html_list(narrative.get('comparison_summary', []))}{self._html_table(payload.get('artifact_comparison', {}).get('metric_deltas', []))}</section>")
            elif section == "article_explanations":
                rows = [{k: v for k, v in item.items() if k in {"article_id", "target", "mode", "prediction", "actual_label", "model", "explanation"}} for item in payload.get("explanations", [])]
                blocks.append(f"<section class=\"report-section\"><h2>Article Explanations</h2>{self._html_table(rows)}</section>")
            elif section == "appendix_filters":
                blocks.append(f"<section class=\"report-section\"><h2>Appendix: Filters</h2>{self._html_table([payload.get('params', {})])}</section>")
            elif section == "appendix_predictions":
                blocks.append(f"<section class=\"report-section\"><h2>Appendix: Predictions</h2><p>Showing first 100 rows in HTML. Excel and JSON retain the full test prediction set.</p>{self._html_table(payload.get('all_predictions', [])[:100])}</section>")
        body = "".join(blocks)
        return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{html.escape(payload.get('report_title') or 'ML Diagnostics Report')}</title>
<style>body{{font-family:Arial,sans-serif;margin:0;background:#f8fafc;color:#0f172a}}main{{max-width:1180px;margin:0 auto;padding:32px}}.cover{{border-left:8px solid #0f766e;background:linear-gradient(135deg,#ffffff,#ecfeff);border-radius:14px;padding:28px;margin-bottom:22px;box-shadow:0 14px 34px rgb(15 23 42 / .08)}}.eyebrow{{color:#0f766e;font-size:12px;font-weight:700;text-transform:uppercase;letter-spacing:.08em}}section{{background:#fff;border:1px solid #dbe3ef;border-radius:10px;padding:20px;margin:18px 0;box-shadow:0 8px 24px rgb(15 23 42 / .05)}}h1{{font-size:34px;margin:10px 0}}h2{{font-size:20px;margin:0 0 14px}}h3{{font-size:15px;margin:18px 0 8px}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{border-bottom:1px solid #e2e8f0;padding:9px;text-align:left;vertical-align:top}}th{{background:#eef6f6;color:#334155;text-transform:uppercase;font-size:11px}}code{{background:#eef2ff;padding:2px 4px;border-radius:4px}}.metrics{{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin-top:18px}}.metric{{border:1px solid #dbe3ef;border-top:4px solid #0f766e;border-radius:8px;padding:12px}}.metric span{{display:block;color:#64748b;font-size:12px}}.metric strong{{display:block;font-size:24px;margin-top:4px}}li{{margin:8px 0}}@page{{size:A4;margin:16mm}}@media print{{body{{background:#fff;color:#0f172a}}main{{max-width:none;padding:0}}.cover,section{{box-shadow:none;border-color:#cbd5e1;break-inside:avoid;page-break-inside:avoid}}.cover{{min-height:45vh;display:flex;flex-direction:column;justify-content:center;break-after:page;page-break-after:always}}table{{font-size:10px;page-break-inside:auto}}thead{{display:table-header-group}}tr{{break-inside:avoid;page-break-inside:avoid}}a{{color:#0f172a;text-decoration:none}}}}</style>
</head><body><main>{body}</main></body></html>"""

    def _compare_html(self, comparison: MLDiagnosticsReportCompareResponse) -> str:
        title = "Diagnostics Report Comparison"
        created = datetime.utcnow().isoformat(timespec="seconds") + "Z"
        body = f"""<header class="cover report-section"><div class="eyebrow">Comparison Report</div><h1>{title}</h1><p>Baseline <code>{html.escape(str(comparison.baseline_report_id or '-'))}</code></p><p>Generated {created}</p></header>"""
        body += f"<section class=\"report-section\"><h2>Compared Reports</h2>{self._html_table([item.model_dump() for item in comparison.items])}</section>"
        body += f"<section class=\"report-section\"><h2>Metrics</h2>{self._html_table(comparison.metrics)}</section>"
        body += f"<section class=\"report-section\"><h2>Metric Deltas</h2>{self._html_table(comparison.metric_deltas)}</section>"
        body += f"<section class=\"report-section\"><h2>Row Count Deltas</h2>{self._html_table(comparison.row_count_deltas)}</section>"
        body += f"<section class=\"report-section\"><h2>Section Matrix</h2>{self._html_table(comparison.section_matrix)}</section>"
        body += f"<section class=\"report-section\"><h2>Failure Summaries</h2>{self._html_table(comparison.failure_summary)}</section>"
        notes = self._html_list(comparison.notes) if comparison.notes else "<p>No comparison notes.</p>"
        body += f"<section class=\"report-section\"><h2>Notes</h2>{notes}</section>"
        return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>{title}</title>
<style>body{{font-family:Arial,sans-serif;margin:0;background:#f8fafc;color:#0f172a}}main{{max-width:1180px;margin:0 auto;padding:32px}}.cover{{border-left:8px solid #0f766e;background:linear-gradient(135deg,#ffffff,#ecfeff);border-radius:14px;padding:28px;margin-bottom:22px;box-shadow:0 14px 34px rgb(15 23 42 / .08)}}.eyebrow{{color:#0f766e;font-size:12px;font-weight:700;text-transform:uppercase;letter-spacing:.08em}}section{{background:#fff;border:1px solid #dbe3ef;border-radius:10px;padding:20px;margin:18px 0;box-shadow:0 8px 24px rgb(15 23 42 / .05)}}h1{{font-size:34px;margin:10px 0}}h2{{font-size:20px;margin:0 0 14px}}table{{border-collapse:collapse;width:100%;font-size:12px}}th,td{{border-bottom:1px solid #e2e8f0;padding:8px;text-align:left;vertical-align:top}}th{{background:#eef6f6;color:#334155;text-transform:uppercase;font-size:10px}}code{{background:#eef2ff;padding:2px 4px;border-radius:4px}}@page{{size:A4;margin:16mm}}@media print{{body{{background:#fff}}main{{max-width:none;padding:0}}.cover,section{{box-shadow:none;border-color:#cbd5e1;break-inside:avoid;page-break-inside:avoid}}.cover{{min-height:42vh;display:flex;flex-direction:column;justify-content:center;break-after:page;page-break-after:always}}thead{{display:table-header-group}}tr{{break-inside:avoid;page-break-inside:avoid}}}}</style>
</head><body><main>{body}</main></body></html>"""

    def _excel(self, path: Path, payload: dict[str, Any]) -> None:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        wb.remove(wb.active)
        diagnostics = payload["diagnostics"]
        narrative = payload.get("narrative", {})
        for section in payload.get("sections", DEFAULT_REPORT_SECTIONS):
            name = ""
            rows = []
            if section == "cover":
                cover = wb.create_sheet("Cover")
                cover["A1"] = payload.get("report_title") or "ML Diagnostics Report"
                cover["A1"].font = Font(bold=True, size=22, color="0F172A")
                cover["A3"] = "Report ID"
                cover["B3"] = payload["report_id"]
                cover["A4"] = "Created"
                cover["B4"] = payload["created_at"]
                cover["A5"] = "Template"
                cover["B5"] = payload.get("template", "portfolio")
                cover["A6"] = "Prepared For"
                cover["B6"] = payload.get("prepared_for") or ""
                cover["A8"] = "Executive Summary"
                cover["B8"] = narrative.get("executive_summary", "")
                cover["B8"].alignment = Alignment(wrap_text=True, vertical="top")
                cover.column_dimensions["A"].width = 18
                cover.column_dimensions["B"].width = 90
                continue
            if section == "executive_summary":
                name = "Executive Summary"
                rows = [{"item": "Executive Summary", "detail": narrative.get("executive_summary", "")}]
                rows.extend({"item": f"Failure {idx + 1}", "detail": line} for idx, line in enumerate(narrative.get("failure_factor_narrative", [])))
                rows.extend({"item": f"Comparison {idx + 1}", "detail": line} for idx, line in enumerate(narrative.get("comparison_summary", [])))
            elif section == "appendix_filters":
                name = "Filters"
                rows = [payload.get("params", {})]
            elif section == "metrics":
                name = "Metrics"
                rows = [diagnostics.get("metrics", {})]
            elif section == "confusion_matrix":
                name = "Confusion Matrix"
                rows = self._matrix_rows(diagnostics.get("labels", []), diagnostics.get("confusion_matrix", []))
            elif section == "class_quality":
                name = "Class Metrics"
                rows = diagnostics.get("class_metrics", [])
            elif section == "failure_narrative":
                name = "Failure Factors"
                rows = diagnostics.get("failure_factors", [])
            elif section == "error_samples":
                name = "Error Samples"
                rows = payload.get("error_samples", [])
            elif section == "appendix_predictions":
                name = "All Predictions"
                rows = payload.get("all_predictions", [])
            elif section == "artifact_comparison":
                name = "Artifact Comparison"
                rows = payload.get("artifact_comparison", {}).get("metric_deltas", [])
            elif section == "article_explanations":
                name = "Explanations"
                rows = [{k: v for k, v in item.items() if k in {"article_id", "target", "mode", "prediction", "actual_label", "model", "explanation"}} for item in payload.get("explanations", [])]
            elif section == "recommendations":
                name = "Recommendations"
                rows = diagnostics.get("recommendations", [])
            if not name:
                continue
            ws = wb.create_sheet(name[:31])
            self._write_sheet(ws, rows, Font(bold=True, color="0F172A"), PatternFill("solid", fgColor="EAF7F6"))
            if name == "Confusion Matrix" and ws.max_row > 1 and ws.max_column > 2:
                end = f"{get_column_letter(ws.max_column)}{ws.max_row}"
                ws.conditional_formatting.add(
                    f"B2:{end}",
                    ColorScaleRule(start_type="min", start_color="FEE2E2", mid_type="percentile", mid_value=50, mid_color="FEF3C7", end_type="max", end_color="86EFAC"),
                )
        wb.save(path)

    def _write_sheet(self, ws: Any, rows: list[dict[str, Any]], header_font: Any, header_fill: Any) -> None:
        from openpyxl.styles import Alignment

        flat_rows = [self._flat_row(row) for row in rows] or [{}]
        headers = sorted({key for row in flat_rows for key in row})
        ws.append(headers or ["value"])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        if headers:
            ws.auto_filter.ref = ws.dimensions
        for row in flat_rows:
            ws.append([row.get(header, "") for header in headers])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if isinstance(cell.value, float) and 0 <= cell.value <= 1:
                    cell.number_format = "0.000"
        for column in ws.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
            ws.column_dimensions[column[0].column_letter].width = width

    def _markdown_table(self, rows: list[dict[str, Any]]) -> str:
        flat_rows = [self._flat_row(row) for row in rows]
        if not flat_rows:
            return "_No data_"
        headers = sorted({key for row in flat_rows for key in row})
        lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
        for row in flat_rows[:100]:
            lines.append("| " + " | ".join(str(row.get(header, "")).replace("\n", " ") for header in headers) + " |")
        return "\n".join(lines)

    def _html_table(self, rows: list[dict[str, Any]]) -> str:
        flat_rows = [self._flat_row(row) for row in rows]
        if not flat_rows:
            return "<p>No data</p>"
        headers = sorted({key for row in flat_rows for key in row})
        head = "".join(f"<th>{html.escape(header)}</th>" for header in headers)
        body = "".join("<tr>" + "".join(f"<td>{html.escape(str(row.get(header, '')))}</td>" for header in headers) + "</tr>" for row in flat_rows[:200])
        return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"

    def _html_list(self, rows: list[str]) -> str:
        if not rows:
            return "<p>No narrative available.</p>"
        return "<ul>" + "".join(f"<li>{html.escape(str(row))}</li>" for row in rows) + "</ul>"

    def _matrix_rows(self, labels: list[str], matrix: list[list[int]]) -> list[dict[str, Any]]:
        return [{"actual": labels[idx], **{label: row[col_idx] for col_idx, label in enumerate(labels)}} for idx, row in enumerate(matrix or []) if idx < len(labels)]

    def _flat_row(self, row: dict[str, Any]) -> dict[str, Any]:
        output = {}
        for key, value in row.items():
            if isinstance(value, (dict, list)):
                output[key] = json.dumps(value, ensure_ascii=False)
            else:
                output[key] = value
        return output

    def _report_dir(self, report_id: str) -> Path:
        safe = report_id.replace("/", "").replace("\\", "")
        return self.reports_dir / safe

    def _read_json(self, path: Path, fallback: Any) -> Any:
        if not path.exists():
            return fallback
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            return fallback
