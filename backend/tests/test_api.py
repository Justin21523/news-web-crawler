from __future__ import annotations

import json
import sys
import time
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]
BACKEND = ROOT / "backend"
for path in (ROOT, BACKEND):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from app.core.config import get_settings  # noqa: E402
from app.main import app  # noqa: E402
from pipeline.db import NewsDB  # noqa: E402
from pipeline.schema import RawArticle  # noqa: E402


def test_health_and_empty_stats(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    with TestClient(app) as client:
        health = client.get("/api/v1/health")
        assert health.status_code == 200
        assert health.json()["status"] == "ok"

        stats = client.get("/api/v1/stats")
        assert stats.status_code == 200
        assert stats.json()["total_articles"] == 0

        llm = client.get("/api/v1/llm/health")
        assert llm.status_code == 200
        assert llm.json()["enabled"] is False
        assert llm.json()["provider"] == "llama.cpp"


def test_articles_empty_list(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    with TestClient(app) as client:
        response = client.get("/api/v1/articles")
        assert response.status_code == 200
        assert response.json()["items"] == []


def test_import_upload_csv_jsonl_and_ingest_job(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    monkeypatch.setenv("NEWS_RAW_DIR", str(tmp_path / "raw"))
    monkeypatch.setenv("NEWS_JOBS_DIR", str(tmp_path / "jobs"))
    get_settings.cache_clear()

    csv_content = "title,content,source,category,publish_date\n測試新聞,這是上傳測試內容,upload,tech,2026-06-25\n"
    with TestClient(app) as client:
        response = client.post(
            "/api/v1/import/upload",
            files={"file": ("articles.csv", csv_content.encode("utf-8"), "text/csv")},
        )
        assert response.status_code == 200
        payload = response.json()
        assert payload["detected_format"] == "csv"
        assert payload["article_count_preview"] == 1
        assert payload["suggested_job"]["type"] == "ingest"
        assert Path(payload["saved_path"]).exists()

        invalid = client.post(
            "/api/v1/import/upload",
            files={"file": ("articles.txt", b"title,content\nx,y\n", "text/plain")},
        )
        assert invalid.status_code == 400

        jsonl = json.dumps({"title": "JSONL", "content": "JSONL 上傳內容"}, ensure_ascii=False) + "\n"
        queued = client.post(
            "/api/v1/import/upload",
            params={"run": "true"},
            files={"file": ("articles.jsonl", jsonl.encode("utf-8"), "application/x-ndjson")},
        )
        assert queued.status_code == 200
        queued_payload = queued.json()
        assert queued_payload["job_id"]
        assert queued_payload["suggested_job"]["params"]["input_dir"]


def test_quality_pipeline_and_analysis_empty_state(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    with TestClient(app) as client:
        quality = client.get("/api/v1/data-quality")
        assert quality.status_code == 200
        assert quality.json()["total_articles"] == 0

        pipeline = client.get("/api/v1/pipeline/status")
        assert pipeline.status_code == 200
        assert len(pipeline.json()["steps"]) == 5

        overview = client.get("/api/v1/analysis/overview")
        assert overview.status_code == 200
        assert overview.json()["nlp_documents"] == 0

        ir = client.get("/api/v1/analysis/ir")
        assert ir.status_code == 200
        assert ir.json()["index_coverage_percent"] == 0


def test_data_quality_returns_article_level_issues(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    db = NewsDB(tmp_path / "news.db")
    db.init()
    raw_path = tmp_path / "raw.jsonl"
    article = RawArticle(
        article_id="quality-1",
        url="not-a-valid-url",
        source="demo",
        source_name="Demo",
        title="短內容品質測試",
        content="內容太短但可被匯入。",
        publish_date="",
    )
    raw_path.write_text(json.dumps(article.to_dict(), ensure_ascii=False) + "\n", encoding="utf-8")
    db.insert_raw_articles(raw_path)
    db.close()

    with TestClient(app) as client:
        quality = client.get("/api/v1/data-quality")
        assert quality.status_code == 200
        payload = quality.json()
        assert payload["quality_score"] < 100
        issue_types = {item["issue_type"] for item in payload["issues"]}
        assert "invalid_url" in issue_types
        assert "short_content" in issue_types


def test_data_quality_issue_filters_export_and_article_detail(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    db = NewsDB(tmp_path / "news.db")
    db.init()
    raw_path = tmp_path / "raw.jsonl"
    rows = [
        RawArticle(
            article_id="quality-a",
            url="bad-url",
            source="demo",
            source_name="Demo",
            title="重複標題測試",
            content="短文。",
            publish_date="",
            category="tech",
            category_name="Technology",
        ),
        RawArticle(
            article_id="quality-b",
            url="https://example.com/quality-b",
            source="other",
            source_name="Other",
            title="重複標題測試",
            content="這是一篇內容長度足夠的新聞資料，用來測試分類、日期與品質問題篩選。" * 3,
            publish_date="2026-06-25",
            category="business",
            category_name="Business",
        ),
    ]
    raw_path.write_text(
        "\n".join(json.dumps(row.to_dict(), ensure_ascii=False) for row in rows) + "\n",
        encoding="utf-8",
    )
    db.insert_raw_articles(raw_path)
    db.close()

    with TestClient(app) as client:
        quality = client.get("/api/v1/data-quality")
        assert quality.status_code == 200
        payload = quality.json()
        assert payload["issue_counts"]["invalid_url"] == 1
        assert "category_coverage" in payload
        assert "length_distribution" in payload

        issues = client.get("/api/v1/data-quality/issues", params={"severity": "danger", "source": "demo"})
        assert issues.status_code == 200
        issue_payload = issues.json()
        assert issue_payload["total"] >= 1
        assert all(item["severity"] == "danger" for item in issue_payload["items"])

        csv_report = client.get("/api/v1/data-quality/export", params={"format": "csv", "issue_type": "invalid_url"})
        assert csv_report.status_code == 200
        assert "text/csv" in csv_report.headers["content-type"]
        assert "quality-a" in csv_report.text

        md_report = client.get("/api/v1/data-quality/export", params={"format": "markdown"})
        assert md_report.status_code == 200
        assert "Data Quality Report" in md_report.text

        article = client.get("/api/v1/articles/quality-a")
        assert article.status_code == 200
        article_issues = {item["issue_type"] for item in article.json()["quality_issues"]}
        assert "invalid_url" in article_issues


def test_analysis_dashboard_endpoints_filters_and_export(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    db = NewsDB(tmp_path / "news.db")
    db.init()
    raw_path = tmp_path / "raw.jsonl"
    rows = [
        RawArticle(
            article_id="analysis-a",
            url="https://example.com/a",
            source="demo",
            source_name="Demo",
            title="AI 資料平台提升新聞分析效率",
            content="人工智慧與資料分析正在協助媒體整理新聞趨勢與來源表現。" * 4,
            publish_date="2026-06-01",
            category="tech",
            category_name="Technology",
        ),
        RawArticle(
            article_id="analysis-b",
            url="https://example.com/b",
            source="other",
            source_name="Other",
            title="商業資料分析協助零售決策",
            content="零售業者利用會員資料、銷售資料與趨勢分析改善補貨策略。" * 4,
            publish_date="2026-06-02",
            category="business",
            category_name="Business",
        ),
    ]
    raw_path.write_text(
        "\n".join(json.dumps(row.to_dict(), ensure_ascii=False) for row in rows) + "\n",
        encoding="utf-8",
    )
    db.insert_raw_articles(raw_path)
    db.close()

    with TestClient(app) as client:
        dashboard = client.get("/api/v1/analysis/dashboard")
        assert dashboard.status_code == 200
        payload = dashboard.json()
        assert payload["kpis"]["total_articles"] == 2
        assert payload["source_distribution"]

        filtered = client.get("/api/v1/analysis/dashboard", params={"source": "demo"})
        assert filtered.status_code == 200
        assert filtered.json()["kpis"]["total_articles"] == 1

        trends = client.get("/api/v1/analysis/trends")
        assert trends.status_code == 200
        assert len(trends.json()["daily_volume"]) == 2

        keywords_entities = client.get("/api/v1/analysis/keywords-entities")
        assert keywords_entities.status_code == 200
        assert keywords_entities.json()["top_entities"]

        csv_report = client.get("/api/v1/analysis/export", params={"section": "overview", "format": "csv"})
        assert csv_report.status_code == 200
        assert "text/csv" in csv_report.headers["content-type"]

        markdown = client.get("/api/v1/analysis/export", params={"section": "business", "format": "markdown"})
        assert markdown.status_code == 200
        assert "Analysis Business" in markdown.text


def test_text_mining_endpoints_on_demo_like_dataset(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    db = NewsDB(tmp_path / "news.db")
    db.init()
    raw_path = tmp_path / "raw.jsonl"
    rows = [
        RawArticle(article_id=f"tm-{i}", url=f"https://example.com/{i}", source="demo" if i % 2 else "other", source_name="Demo", title=f"資料分析新聞 {i}", content=("資料 分析 人工智慧 新聞 趨勢 商業 來源 " * 8), publish_date=f"2026-06-{i:02d}", category="tech" if i % 2 else "business", category_name="Technology" if i % 2 else "Business")
        for i in range(1, 8)
    ]
    raw_path.write_text("\n".join(json.dumps(row.to_dict(), ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")
    db.insert_raw_articles(raw_path)
    from nlp.pipeline import NLPPipeline
    from pipeline.schema import EnrichedArticle
    nlp = NLPPipeline(engine="jieba")
    for row in rows:
        result = nlp.process(f"{row.title} {row.content}")
        db.upsert_nlp(EnrichedArticle(article_id=row.article_id, url=row.url, source=row.source, title=row.title, content_clean=row.content, publish_date=row.publish_date, category=row.category, category_name=row.category_name, tags=[], tokens=result.tokens, keywords=result.keywords, entities=[]))
    db.close()

    with TestClient(app) as client:
        overview = client.get("/api/v1/text-mining/overview")
        assert overview.status_code == 200
        assert overview.json()["total_documents"] == 7
        assert overview.json()["top_keywords"]

        tfidf = client.get("/api/v1/text-mining/tfidf")
        assert tfidf.status_code == 200
        assert tfidf.json()["vocabulary_size"] > 0

        ngrams = client.get("/api/v1/text-mining/ngrams", params={"n": 2})
        assert ngrams.status_code == 200
        assert ngrams.json()["ngrams"]

        topics = client.get("/api/v1/text-mining/topics", params={"n_topics": 3})
        assert topics.status_code == 200
        assert topics.json()["status"] in {"ready", "insufficient_data", "error"}

        clusters = client.get("/api/v1/text-mining/clusters", params={"n_clusters": 3})
        assert clusters.status_code == 200
        assert clusters.json()["status"] in {"ready", "insufficient_data", "error"}

        similarity = client.get("/api/v1/text-mining/similarity", params={"query": "資料 分析"})
        assert similarity.status_code == 200
        assert similarity.json()["results"]

        network = client.get("/api/v1/text-mining/network", params={"type": "entity"})
        assert network.status_code == 200
        assert network.json()["nodes"]

        collocations = client.get("/api/v1/text-mining/collocations", params={"metric": "npmi", "level": "document"})
        assert collocations.status_code == 200
        assert "collocations" in collocations.json()

        co_occurrence = client.get("/api/v1/text-mining/co-occurrence", params={"kind": "keyword", "weight_metric": "npmi"})
        assert co_occurrence.status_code == 200
        assert "nodes" in co_occurrence.json()

        relationships = client.get("/api/v1/text-mining/relationships")
        assert relationships.status_code == 200
        assert "share_of_voice" in relationships.json()

        bursts = client.get("/api/v1/text-mining/bursts", params={"recent_window_days": 3, "baseline_window_days": 6})
        assert bursts.status_code == 200
        assert "rising_terms" in bursts.json()

        entities = client.get("/api/v1/text-mining/entities")
        assert entities.status_code == 200
        assert entities.json()["entity_frequency"]
        assert entities.json()["representative_articles"]
        assert entities.json()["active_entity_source"]

        providers = client.get("/api/v1/entity-extraction/providers")
        assert providers.status_code == 200
        assert any(item["provider"] == "fallback" and item["available"] for item in providers.json()["providers"])

        run_a = client.post("/api/v1/entity-extraction/run", params={"provider": "fallback", "activate": True})
        assert run_a.status_code == 200
        assert run_a.json()["run"]["provider"] == "fallback"
        run_b = client.post("/api/v1/entity-extraction/run", params={"provider": "fallback", "activate": False})
        assert run_b.status_code == 200

        compare = client.get("/api/v1/entity-extraction/compare")
        assert compare.status_code == 200
        assert compare.json()["status"] == "ready"
        assert compare.json()["metrics"]["overlap_articles"] > 0

        export = client.get("/api/v1/text-mining/export", params={"section": "ngrams", "format": "csv"})
        assert export.status_code == 200
        assert "text/csv" in export.headers["content-type"]

        collocation_export = client.get("/api/v1/text-mining/export", params={"section": "collocations", "format": "json"})
        assert collocation_export.status_code == 200
        assert "application/json" in collocation_export.headers["content-type"]

        assignments = client.post("/api/v1/text-mining/assignments", params={"assignment_type": "both", "n_topics": 3, "n_clusters": 3})
        assert assignments.status_code == 200
        assignment_payload = assignments.json()
        assert assignment_payload["run_id"]
        assert assignment_payload["assignment_counts"]

        latest = client.get("/api/v1/text-mining/assignments/latest")
        assert latest.status_code == 200
        assert latest.json()["runs"]

        facets = client.get("/api/v1/facets")
        assert facets.status_code == 200
        topic_counts = facets.json()["topic_counts"]
        cluster_counts = facets.json()["cluster_counts"]
        assert topic_counts or cluster_counts

        if topic_counts:
            filtered = client.get("/api/v1/articles", params={"topic": topic_counts[0]["value"]})
            assert filtered.status_code == 200
            assert filtered.json()["total"] >= 1

        detail = client.get("/api/v1/articles/tm-1")
        assert detail.status_code == 200
        assert "assignments" in detail.json()


def test_ml_baselines_decision_tree_and_exports(tmp_path, monkeypatch):
    monkeypatch.setenv("NEWS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("NEWS_DB_PATH", str(tmp_path / "news.db"))
    get_settings.cache_clear()

    db = NewsDB(tmp_path / "news.db")
    db.init()
    raw_path = tmp_path / "raw.jsonl"
    rows = []
    for i in range(1, 13):
        category = "Technology" if i <= 6 else "Business"
        source = ["cna", "pts", "ltn"][i % 3]
        sentiment_terms = "success growth improve opportunity" if i % 2 else "risk crisis problem loss"
        rows.append(
            RawArticle(
                article_id=f"ml-{i}",
                url=f"https://example.com/ml-{i}",
                source=source,
                source_name=source.upper(),
                title=f"{category} source {source} report {i}",
                content=(f"{category} analysis {source} {sentiment_terms} market data model insight " * 8),
                publish_date=f"2026-06-{i:02d}",
                category=category.lower(),
                category_name=category,
            )
        )
    raw_path.write_text("\n".join(json.dumps(row.to_dict(), ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")
    db.insert_raw_articles(raw_path)
    db.close()

    with TestClient(app) as client:
        overview = client.get("/api/v1/ml/overview")
        assert overview.status_code == 200
        assert overview.json()["targets"]["source"]["status"] == "ready"

        dataset = client.get("/api/v1/ml/dataset", params={"target": "sentiment"})
        assert dataset.status_code == 200
        assert dataset.json()["label_distribution"]
        assert "weak labels" in " ".join(dataset.json()["notes"])

        source_model = client.get("/api/v1/ml/train", params={"target": "source", "model": "logistic_regression", "feature_limit": 500})
        assert source_model.status_code == 200
        source_payload = source_model.json()
        assert source_payload["status"] == "ready"
        assert "accuracy" in source_payload["metrics"]
        assert source_payload["confusion_matrix"]

        diagnostics = client.get("/api/v1/ml/diagnostics", params={"target": "source", "model": "logistic_regression", "feature_limit": 500})
        assert diagnostics.status_code == 200
        diagnostics_payload = diagnostics.json()
        assert diagnostics_payload["status"] == "ready"
        assert "failure_factors" in diagnostics_payload
        assert "recommendations" in diagnostics_payload

        errors = client.get("/api/v1/ml/diagnostics/errors", params={"target": "source", "model": "logistic_regression", "feature_limit": 500, "page_size": 5})
        assert errors.status_code == 200
        assert "items" in errors.json()

        prediction = client.get("/api/v1/ml/predict/article/ml-1", params={"targets": "category,source,sentiment", "prefer_artifact": False})
        assert prediction.status_code == 200
        prediction_payload = prediction.json()
        assert prediction_payload["article_id"] == "ml-1"
        assert prediction_payload["predictions"]

        linear_explanation = client.get("/api/v1/ml/predict/article/ml-1/explanation", params={"target": "source", "mode": "linear_coefficients", "feature_limit": 500})
        assert linear_explanation.status_code == 200
        assert linear_explanation.json()["mode"] in {"linear_coefficients", "top_terms"}
        assert linear_explanation.json()["contributions"]

        compare = client.get("/api/v1/ml/compare", params={"target": "category", "feature_limit": 500})
        assert compare.status_code == 200
        assert len(compare.json()["results"]) == 5

        tree = client.get("/api/v1/ml/decision-tree", params={"target": "category", "max_depth": 3, "feature_limit": 500})
        assert tree.status_code == 200
        tree_payload = tree.json()
        assert tree_payload["status"] == "ready"
        assert tree_payload["tree"]["node_count"] >= 1
        assert tree_payload["feature_importance"]

        path = client.get("/api/v1/ml/decision-tree/path", params={"target": "category", "article_id": "ml-1", "max_depth": 3, "feature_limit": 500})
        assert path.status_code == 200
        assert path.json()["prediction"] in {"Technology", "Business"}
        assert path.json()["path"]

        svg = client.get("/api/v1/ml/decision-tree/image", params={"target": "category", "format": "svg", "max_depth": 3, "feature_limit": 500})
        assert svg.status_code == 200
        assert "image/svg+xml" in svg.headers["content-type"]
        assert "<svg" in svg.text

        feature_csv = client.get("/api/v1/ml/export", params={"target": "category", "model": "decision_tree", "section": "feature_importance", "format": "csv", "feature_limit": 500})
        assert feature_csv.status_code == 200
        assert "text/csv" in feature_csv.headers["content-type"]
        assert "feature" in feature_csv.text

        job = client.post(
            "/api/v1/jobs",
            json={"type": "train_decision_tree", "params": {"target": "category", "max_depth": 3, "feature_limit": 500}},
        )
        assert job.status_code == 201
        job_id = job.json()["id"]
        final_job = job.json()
        for _ in range(40):
            polled = client.get(f"/api/v1/jobs/{job_id}")
            assert polled.status_code == 200
            final_job = polled.json()
            if final_job["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.25)
        assert final_job["status"] == "succeeded", client.get(f"/api/v1/jobs/{job_id}/logs").json()["content"]

        artifacts = client.get("/api/v1/ml/artifacts")
        assert artifacts.status_code == 200
        items = artifacts.json()["items"]
        assert items
        artifact = items[0]
        assert artifact["job_id"] == job_id
        assert artifact["status"] == "ready"
        assert "model.joblib" in artifact["files"]
        assert "vectorizer.joblib" in artifact["files"]
        assert "decision_tree.svg" in artifact["files"]
        assert "diagnostics.json" in artifact["files"]

        detail = client.get(f"/api/v1/ml/artifacts/{artifact['artifact_id']}")
        assert detail.status_code == 200
        assert detail.json()["manifest"]["artifact_id"] == artifact["artifact_id"]
        assert detail.json()["manifest"]["diagnostics_version"] == "1"

        compare_artifacts = client.get("/api/v1/ml/artifacts/compare", params={"ids": artifact["artifact_id"]})
        assert compare_artifacts.status_code == 200
        assert compare_artifacts.json()["items"][0]["artifact_id"] == artifact["artifact_id"]

        diagnostics_compare = client.get("/api/v1/ml/artifacts/diagnostics/compare", params={"target": "category", "ids": artifact["artifact_id"]})
        assert diagnostics_compare.status_code == 200
        assert diagnostics_compare.json()["items"][0]["artifact_id"] == artifact["artifact_id"]

        artifact_errors = client.get("/api/v1/ml/diagnostics/errors", params={"target": "category", "artifact_id": artifact["artifact_id"]})
        assert artifact_errors.status_code == 200
        assert "items" in artifact_errors.json()

        tree_explanation = client.get("/api/v1/ml/predict/article/ml-1/explanation", params={"target": "category", "mode": "tree_path", "artifact_id": artifact["artifact_id"]})
        assert tree_explanation.status_code == 200
        assert tree_explanation.json()["mode"] == "tree_path"
        assert tree_explanation.json()["path"]

        model_download = client.get(f"/api/v1/ml/artifacts/{artifact['artifact_id']}/download", params={"file": "model"})
        assert model_download.status_code == 200
        assert model_download.content

        artifact_prediction = client.get("/api/v1/ml/predict/article/ml-1", params={"targets": "category", "prefer_artifact": True})
        assert artifact_prediction.status_code == 200
        assert artifact_prediction.json()["predictions"][0]["model_source"] in {"artifact", "on_demand"}

        bad_download = client.get(f"/api/v1/ml/artifacts/{artifact['artifact_id']}/download", params={"file": "../news.db"})
        assert bad_download.status_code == 400

        report_job = client.post(
            "/api/v1/jobs",
            json={
                "type": "export_ml_diagnostics_report",
                "params": {
                    "target": "source",
                    "model": "logistic_regression",
                    "artifact_ids": [artifact["artifact_id"]],
                    "article_ids": ["ml-1"],
                    "feature_limit": 500,
                    "max_depth": 3,
                    "limit": 2000,
                    "template": "portfolio",
                    "sections": ["cover", "metrics", "executive_summary", "artifact_comparison", "failure_narrative", "recommendations", "appendix_predictions"],
                    "report_title": "Portfolio ML Diagnostics",
                    "prepared_for": "Hiring Review",
                },
            },
        )
        assert report_job.status_code == 201
        report_job_id = report_job.json()["id"]
        final_report_job = report_job.json()
        for _ in range(40):
            polled = client.get(f"/api/v1/jobs/{report_job_id}")
            assert polled.status_code == 200
            final_report_job = polled.json()
            if final_report_job["status"] in {"succeeded", "failed"}:
                break
            time.sleep(0.25)
        assert final_report_job["status"] == "succeeded", client.get(f"/api/v1/jobs/{report_job_id}/logs").json()["content"]

        reports = client.get("/api/v1/ml/diagnostics/reports")
        assert reports.status_code == 200
        report_items = reports.json()["items"]
        assert report_items
        latest_report = report_items[0]
        assert latest_report["job_id"] == report_job_id
        assert latest_report["template"] == "portfolio"
        assert latest_report["report_title"] == "Portfolio ML Diagnostics"
        assert latest_report["prepared_for"] == "Hiring Review"
        assert "failure_narrative" in latest_report["sections"]
        assert latest_report["sections"][:3] == ["cover", "metrics", "executive_summary"]
        assert latest_report["section_status"]["metrics"] == "ready"
        assert latest_report["preview_url"].endswith("/download?format=html")
        assert latest_report["print_url"].endswith("/download?format=html")
        assert latest_report["pdf_url"].endswith("/download?format=pdf")
        assert latest_report["pdf_status"] in {"not_generated", "ready", "fallback_html", "error"}
        assert {"report.html", "report.md", "report.xlsx", "report.json"}.issubset(set(latest_report["files"]))
        assert latest_report["row_counts"]["all_predictions"] >= 1

        filtered_reports = client.get("/api/v1/ml/diagnostics/reports", params={"q": "Portfolio", "target": "source", "model": "logistic_regression", "sort_by": "f1_macro", "page": 1, "page_size": 5})
        assert filtered_reports.status_code == 200
        assert filtered_reports.json()["total"] >= 1

        report_detail = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}")
        assert report_detail.status_code == 200
        assert report_detail.json()["diagnostics"]["status"] == "ready"

        html_report = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/download", params={"format": "html"})
        assert html_report.status_code == 200
        assert "text/html" in html_report.headers["content-type"]
        assert "Portfolio ML Diagnostics" in html_report.text
        assert "Executive Summary" in html_report.text
        assert "Failure Factor Narrative" in html_report.text
        assert "@media print" in html_report.text
        assert "@page" in html_report.text
        assert html_report.text.index("<h2>Metrics</h2>") < html_report.text.index("<h2>Executive Summary</h2>")

        markdown_report = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/download", params={"format": "markdown"})
        assert markdown_report.status_code == 200
        assert "Portfolio ML Diagnostics" in markdown_report.text
        assert "## Failure Factor Narrative" in markdown_report.text
        assert markdown_report.text.index("## Metrics") < markdown_report.text.index("## Executive Summary")

        pdf_report = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/download", params={"format": "pdf"})
        assert pdf_report.status_code == 200
        if "application/pdf" in pdf_report.headers.get("content-type", ""):
            assert pdf_report.content.startswith(b"%PDF")
        else:
            assert "text/html" in pdf_report.headers["content-type"]
            assert pdf_report.headers["x-report-fallback"] == "html"
            assert "Portfolio ML Diagnostics" in pdf_report.text

        excel_report = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/download", params={"format": "xlsx"})
        assert excel_report.status_code == 200
        assert excel_report.content.startswith(b"PK")
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(excel_report.content), read_only=False)
        assert {"Cover", "Executive Summary", "Metrics", "Failure Factors", "Artifact Comparison", "All Predictions"}.issubset(set(workbook.sheetnames))
        assert workbook.sheetnames[:3] == ["Cover", "Metrics", "Executive Summary"]

        json_report = client.get(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/download", params={"format": "json"})
        assert json_report.status_code == 200
        json_payload = json_report.json()
        assert json_payload["report_id"] == latest_report["report_id"]
        assert json_payload["template"] == "portfolio"
        assert json_payload["narrative"]["executive_summary"]

        report_compare = client.get("/api/v1/ml/diagnostics/reports/compare", params={"ids": f"{latest_report['report_id']},{latest_report['report_id']}"})
        assert report_compare.status_code == 200
        compare_payload = report_compare.json()
        assert compare_payload["baseline_report_id"] == latest_report["report_id"]
        assert compare_payload["metrics"]
        assert compare_payload["metric_deltas"]
        assert compare_payload["section_matrix"]

        metadata_update = client.patch(f"/api/v1/ml/diagnostics/reports/{latest_report['report_id']}/metadata", json={"tags": ["Portfolio", "Review"]})
        assert metadata_update.status_code == 200
        assert metadata_update.json()["tags"] == ["portfolio", "review"]

        tagged_reports = client.get("/api/v1/ml/diagnostics/reports", params={"tag": "portfolio"})
        assert tagged_reports.status_code == 200
        assert tagged_reports.json()["total"] >= 1

        archive = client.post("/api/v1/ml/diagnostics/reports/bulk", json={"action": "archive", "ids": [latest_report["report_id"]]})
        assert archive.status_code == 200
        assert archive.json()["updated"] == 1
        archived = client.get("/api/v1/ml/diagnostics/reports", params={"status": "archived"})
        assert archived.status_code == 200
        assert any(item["report_id"] == latest_report["report_id"] for item in archived.json()["items"])

        trash = client.post("/api/v1/ml/diagnostics/reports/bulk", json={"action": "trash", "ids": [latest_report["report_id"]]})
        assert trash.status_code == 200
        hidden = client.get("/api/v1/ml/diagnostics/reports")
        assert all(item["report_id"] != latest_report["report_id"] for item in hidden.json()["items"])
        trashed = client.get("/api/v1/ml/diagnostics/reports", params={"status": "trashed", "include_trashed": True})
        assert any(item["report_id"] == latest_report["report_id"] for item in trashed.json()["items"])

        restore = client.post("/api/v1/ml/diagnostics/reports/bulk", json={"action": "restore", "ids": [latest_report["report_id"]]})
        assert restore.status_code == 200
        assert restore.json()["items"][0]["status"] == "active"

        bulk_zip = client.post("/api/v1/ml/diagnostics/reports/bulk/export", json={"ids": [latest_report["report_id"]], "formats": ["html", "json", "manifest"]})
        assert bulk_zip.status_code == 200
        assert "application/zip" in bulk_zip.headers["content-type"]
        assert bulk_zip.content.startswith(b"PK")

        compare_html = client.get("/api/v1/ml/diagnostics/reports/compare/download", params={"ids": f"{latest_report['report_id']},{latest_report['report_id']}", "format": "html"})
        assert compare_html.status_code == 200
        assert "Diagnostics Report Comparison" in compare_html.text

        compare_pdf = client.get("/api/v1/ml/diagnostics/reports/compare/download", params={"ids": f"{latest_report['report_id']},{latest_report['report_id']}", "format": "pdf"})
        assert compare_pdf.status_code == 200
        assert "application/pdf" in compare_pdf.headers.get("content-type", "") or compare_pdf.headers.get("x-report-fallback") == "html"
